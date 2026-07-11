from __future__ import annotations

import re
import json
import shutil
from collections import defaultdict
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOC = Path("/Users/chengwen/Downloads/黑曜纪元_种族与势力设定文档_v0.1.docx")
IMAGE_CARD = Path("/Users/chengwen/Downloads/ChatGPT Image 2026年6月4日 18_06_23.png")
IMAGE_CONCEPT = Path("/Users/chengwen/Downloads/ChatGPT Image 2026年6月4日 18_06_31.png")

FOLDERS = {
    "overview": "00_总览",
    "world": "01_世界底层",
    "races": "02_种族",
    "factions": "03_势力组织",
    "creatures": "04_异化生物",
    "places": "05_地点生态",
    "plot": "06_事件剧情",
    "bases": "07_基底气质",
    "relations": "08_关系网络",
    "assets": "09_素材与图片",
    "templates": "_templates",
    "docs": "docs/superpowers/specs",
}


def mkdirs() -> None:
    for folder in FOLDERS.values():
        (ROOT / folder).mkdir(parents=True, exist_ok=True)
    (ROOT / ".obsidian").mkdir(parents=True, exist_ok=True)


def clean_filename(name: str) -> str:
    name = name.strip()
    name = re.sub(r'[\\/:*?"<>|]', "／", name)
    name = re.sub(r"\s+", " ", name)
    return name


def yaml_value(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(f"  - {item}" for item in value)
    text = str(value).replace('"', '\\"')
    return f'"{text}"'


def frontmatter(**fields) -> str:
    lines = ["---"]
    for key, value in fields.items():
        if isinstance(value, list):
            lines.append(f"{key}:")
            if value:
                lines.extend(f"  - {item}" for item in value)
            else:
                lines.append("  - ")
        else:
            lines.append(f"{key}: {yaml_value(value)}")
    lines.append("---")
    return "\n".join(lines) + "\n\n"


def wikilink(name: str) -> str:
    display = name.strip()
    target = clean_filename(display)
    if target != display:
        return f"[[{target}|{display}]]"
    return f"[[{target}]]"


def link_list(items: list[str]) -> str:
    return "、".join(wikilink(item) for item in items if item)


def split_list(text: str) -> list[str]:
    if not text:
        return []
    parts = re.split(r"\s*[、,/／]\s*|\s+/\s+", text.strip())
    return [part.strip() for part in parts if part.strip()]


def write_note(folder_key: str, filename: str, body: str) -> Path:
    path = ROOT / FOLDERS[folder_key] / f"{clean_filename(filename)}.md"
    path.write_text(body, encoding="utf-8")
    return path


def parse_keyed_lines(text: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    for line in [item.strip() for item in text.splitlines() if item.strip()]:
        if "：" in line:
            key, value = line.split("：", 1)
            fields[key.strip()] = value.strip()
    return fields


def extract_card_table(table) -> dict[str, object]:
    row = table.rows[0]
    left = [line.strip() for line in row.cells[0].text.splitlines() if line.strip()]
    right = parse_keyed_lines(row.cells[1].text)
    name = left[0]
    category = left[1] if len(left) > 1 else ""
    base_line = ""
    tagline = ""
    for line in left[2:]:
        if line.startswith("基底："):
            base_line = line.removeprefix("基底：").strip()
        else:
            tagline = line
    bases = split_list(base_line)
    return {
        "name": name,
        "category": category,
        "bases": bases,
        "tagline": tagline,
        "appearance": right.get("外观/气质", ""),
        "ability": right.get("核心能力", ""),
        "cost": right.get("代价/弱点", ""),
        "story": right.get("剧情用途", ""),
    }


def classify_folder(category: str) -> str:
    if "异化生物" in category:
        return "creatures"
    if any(word in category for word in ["组织", "宗门", "机构", "联盟", "部族", "神使", "群体", "秘社", "武装"]):
        return "factions"
    return "races"


def make_entity_note(entity: dict[str, object], source_section: str) -> str:
    name = str(entity["name"])
    category = str(entity.get("category", ""))
    bases = list(entity.get("bases", []))
    tags = ["黑曜纪元", "设定条目"]
    if "异化生物" in category:
        tags.append("异化生物")
    elif classify_folder(category) == "factions":
        tags.append("势力组织")
    else:
        tags.append("种族")

    body = frontmatter(
        type=category or "设定条目",
        base=bases,
        status="草案",
        source="黑曜纪元_种族与势力设定文档_v0.1.docx",
        source_section=source_section,
        tags=tags,
    )
    body += f"# {name}\n\n"
    if entity.get("tagline"):
        body += f"> {entity['tagline']}\n\n"
    body += "## 关系\n"
    body += f"- 分类：{category or '未分类'}\n"
    if bases:
        body += f"- 基底：{link_list(bases)}\n"
    body += f"- 来源章节：{wikilink(source_section)}\n\n"

    fields = [
        ("外观与气质", entity.get("appearance", "")),
        ("核心能力", entity.get("ability", "")),
        ("代价与弱点", entity.get("cost", "")),
        ("剧情用途", entity.get("story", "")),
    ]
    for title, value in fields:
        if value:
            body += f"## {title}\n{value}\n\n"
    body += "## 待补完\n- 语言 / 交流方式\n- 社会结构或群落结构\n- 繁衍方式\n- 代表角色或遭遇事件\n"
    return body


def make_base_notes(table) -> list[dict[str, str]]:
    bases = []
    for row in table.rows[1:]:
        cells = [cell.text.strip() for cell in row.cells]
        base = {
            "name": cells[0],
            "mood": cells[1],
            "colors": cells[2],
            "fits": cells[3],
        }
        bases.append(base)
        fit_items = split_list(cells[3].replace("、", "/"))
        body = frontmatter(
            type="基底气质",
            status="草案",
            source="黑曜纪元_种族与势力设定文档_v0.1.docx",
            tags=["黑曜纪元", "基底气质"],
        )
        body += f"# {base['name']}\n\n"
        body += f"> {base['mood']}\n\n"
        body += f"## 色彩方向\n{base['colors']}\n\n"
        body += "## 适配对象\n"
        for item in fit_items:
            body += f"- {wikilink(item)}\n"
        body += "\n## 使用提醒\n该基底决定卡片的情绪底色；阵营身份应通过识别色、纹章、材质和 UI 细节叠加，而不是完全替代基底。\n"
        write_note("bases", base["name"], body)
    return bases


def make_creature_ecology_notes(tables) -> list[dict[str, object]]:
    section_names = [
        "10.1 城市异化生物",
        "10.2 深海异化生物",
        "10.3 森林 / 山海异化生物",
        "10.4 神尸矿区异化生物",
        "10.5 星外异化生物",
    ]
    creatures: list[dict[str, object]] = []
    for table, section in zip(tables, section_names, strict=True):
        for row in table.rows[1:]:
            name, niche, setting, base_text = [cell.text.strip() for cell in row.cells]
            bases = split_list(base_text)
            creature = {
                "name": name,
                "category": "异化生物",
                "niche": niche,
                "setting": setting,
                "bases": bases,
                "source_section": section,
            }
            creatures.append(creature)
            body = frontmatter(
                type="异化生物",
                base=bases,
                ecology=niche,
                status="草案",
                source="黑曜纪元_种族与势力设定文档_v0.1.docx",
                source_section=section,
                tags=["黑曜纪元", "异化生物", "生态条目"],
            )
            body += f"# {name}\n\n"
            body += f"> {setting}\n\n"
            body += "## 关系\n"
            body += f"- 生态位：{wikilink(niche)}\n"
            if bases:
                body += f"- 推荐基底：{link_list(bases)}\n"
            body += f"- 来源章节：{wikilink(section)}\n\n"
            body += "## 可扩展方向\n- 栖息地\n- 捕食方式\n- 可采集资源\n- 遭遇事件\n- 弱点与克制方式\n"
            write_note("creatures", name, body)
    return creatures


def copy_assets() -> None:
    shutil.copy2(IMAGE_CARD, ROOT / FOLDERS["assets"] / "孢雾巡猎者_档案卡.png")
    shutil.copy2(IMAGE_CONCEPT, ROOT / FOLDERS["assets"] / "孢雾巡猎者_概念展示.png")


def make_spore_hunter_note() -> None:
    body = frontmatter(
        type="异化生物",
        subtype=["腐生变异体", "掠食型"],
        base=["雾绿腐生", "深黑不可名状"],
        habitat=["腐林", "湿谷", "废弃矿区"],
        threat="C",
        weakness=["火焰", "强光", "干燥环境"],
        status="图像设定草案",
        source="ChatGPT Image 2026年6月4日 18:06",
        tags=["黑曜纪元", "异化生物", "视觉设定", "孢雾"],
    )
    body += "# 孢雾巡猎者\n\n"
    body += "> 腐林、湿谷与废弃矿区中的群猎型腐生变异体；单体威胁有限，群体活动时危险上升。\n\n"
    body += "## 图像素材\n"
    body += "![[孢雾巡猎者_概念展示.png]]\n\n"
    body += "![[孢雾巡猎者_档案卡.png]]\n\n"
    body += "## 关系\n"
    body += f"- 分类：{wikilink('异化生物生态')} / {wikilink('腐生变异体')} / {wikilink('掠食型')}\n"
    body += f"- 基底：{wikilink('雾绿腐生')} / {wikilink('深黑不可名状')}\n"
    body += f"- 栖息地：{wikilink('腐林')}、{wikilink('湿谷')}、{wikilink('废弃矿区')}\n"
    body += f"- 可能冲突：{wikilink('赛博工业财团')}、{wikilink('管巢族')}、{wikilink('神尸矿区')}\n"
    body += f"- 可被利用：背部孢囊可被采集为{wikilink('生物武器材料')}或{wikilink('腐化孢子样本')}\n\n"
    body += "## 外观\n"
    body += "头部覆盖骨质面甲，胸背部呈黑色湿亮甲壳，背脊生有半透明孢囊；四肢修长，利爪用于撕裂装甲与厚重组织，尾部延展出湿膜鳍片，适合在湿地快速变向。\n\n"
    body += "## 能力\n"
    body += "- 孢雾遮蔽：释放麻痹性腐雾，干扰猎物感官。\n"
    body += "- 孢囊储存：背部孢囊储存腐化孢子，受损时会释放致幻与腐蚀性雾气。\n"
    body += "- 骨质面甲：前颅骨质硬壳能抵御多数物理冲击。\n"
    body += "- 湿膜尾鳍：在湿地、浅水和腐泥环境中高速滑行。\n\n"
    body += "## 弱点\n"
    body += f"{wikilink('火焰')}、{wikilink('强光')}与{wikilink('干燥环境')}会破坏其孢囊活性并降低行动能力。\n\n"
    body += "## 设定校正\n"
    body += "图中标题若写作“抱雾巡猎者”，建议统一改为“孢雾巡猎者”；背部结构也建议使用“孢囊背脊”，以匹配腐生、菌雾与孢子生态。\n"
    write_note("creatures", "孢雾巡猎者", body)


def make_world_notes(doc: Document, system_table) -> None:
    core_logic = doc.paragraphs[23].text.strip()
    systems = []
    for row in system_table.rows[1:]:
        system, explanation, function = [cell.text.strip() for cell in row.cells]
        systems.append((system, explanation, function))

    body = frontmatter(
        type="世界观总览",
        status="草案",
        source="黑曜纪元_种族与势力设定文档_v0.1.docx",
        tags=["黑曜纪元", "世界观"],
    )
    body += "# 黑曜纪元\n\n"
    body += f"{core_logic}\n\n"
    body += "## 核心原则\n"
    body += "种族不是“换皮职业”，而是世界裂灾后不同文明、血脉、技术和污染方式形成的生存形态。\n\n"
    body += "## 关键入口\n"
    body += "- [[诸天裂灾]]\n- [[视觉基底索引]]\n- [[种族与势力索引]]\n- [[冲突关系索引]]\n- [[异化生物生态]]\n"
    write_note("world", "黑曜纪元", body)

    body = frontmatter(
        type="世界事件",
        status="草案",
        source="黑曜纪元_种族与势力设定文档_v0.1.docx",
        tags=["黑曜纪元", "世界事件"],
    )
    body += "# 诸天裂灾\n\n"
    body += "多个维度、神系、梦境层和星外逻辑同时撕开现实。不同文明用不同术语解释同一场事故，由此形成科技、修行、魔法、神系与旧神眷族并存的格局。\n\n"
    body += "## 各体系解释\n"
    body += "| 体系 | 解释方式 | 实际作用 |\n| --- | --- | --- |\n"
    for system, explanation, function in systems:
        body += f"| {wikilink(system)} | {explanation} | {function} |\n"
    body += f"\n## 反向连接\n- {wikilink('黑曜纪元')}\n- {wikilink('主要冲突关系')}\n"
    write_note("world", "诸天裂灾", body)


def make_hub_notes() -> None:
    hubs = {
        "异化生物生态": ("plot", "异化生物用于填充野外、城市、深海、矿区和星外区域。它们既是怪物，也是资源、灾害、仪式材料和剧情线索。"),
        "主要冲突关系": ("relations", "跨体系冲突构成世界推进的主引擎。"),
        "腐林": ("places", "腐生孢子覆盖的森林生态，适合雾绿腐生基底。"),
        "湿谷": ("places", "湿冷、低光、孢雾沉积的谷地生态。"),
        "废弃矿区": ("places", "工业废弃与神秘污染交叠的矿区。"),
        "神尸矿区": ("places", "神尸矿脉、工业采掘和异化生物共同形成的危险矿区；异化生态是矿区内部状态，不是独立地点。"),
        "赛博工业财团": ("factions", "将神秘力量拆解为供应链、义体、能源和商品的工业权力。"),
    }
    for name, (folder, summary) in hubs.items():
        body = frontmatter(type="索引节点", status="草案", tags=["黑曜纪元", "图谱节点"])
        body += f"# {name}\n\n{summary}\n\n## 关联\n- {wikilink('黑曜纪元')}\n"
        write_note(folder, name, body)


def make_relation_notes(table) -> list[dict[str, str]]:
    relations = []
    for row in table.rows[1:]:
        axis, summary = [cell.text.strip() for cell in row.cells]
        sides = [part.strip() for part in axis.split(" vs ")]
        filename = axis.replace(" vs ", "_vs_")
        relations.append({"axis": axis, "summary": summary, "filename": filename})
        body = frontmatter(
            type="冲突轴",
            sides=sides,
            status="草案",
            source="黑曜纪元_种族与势力设定文档_v0.1.docx",
            tags=["黑曜纪元", "冲突关系"],
        )
        body += f"# {axis}\n\n{summary}\n\n"
        body += "## 双方\n"
        for side in sides:
            body += f"- {wikilink(side)}\n"
        body += "\n## 剧情用途\n- 阵营任务\n- 副本冲突\n- 主线价值观对撞\n"
        write_note("relations", filename, body)
    return relations


def make_merge_and_priority_notes(merge_table, priority_table) -> None:
    body = frontmatter(type="整理建议", status="草案", tags=["黑曜纪元", "设定整理"])
    body += "# 合并建议\n\n"
    body += "| 当前概念 | 建议处理 | 理由 |\n| --- | --- | --- |\n"
    for row in merge_table.rows[1:]:
        concept, handling, reason = [cell.text.strip() for cell in row.cells]
        body += f"| {concept} | {handling} | {reason} |\n"
    write_note("relations", "合并建议", body)

    body = frontmatter(type="主轴规划", status="草案", tags=["黑曜纪元", "主轴"])
    body += "# 优先主轴\n\n"
    body += "| 优先级 | 保留对象 | 功能 |\n| --- | --- | --- |\n"
    for row in priority_table.rows[1:]:
        priority, objects, function = [cell.text.strip() for cell in row.cells]
        linked = "、".join(wikilink(item.strip()) for item in re.split(r"、|,", objects) if item.strip())
        body += f"| {priority} | {linked} | {function} |\n"
    write_note("relations", "优先主轴", body)


def make_indexes(entities, bases, creatures, relations) -> None:
    grouped = defaultdict(list)
    for entity in entities:
        grouped[classify_folder(str(entity["category"]))].append(str(entity["name"]))
    for creature in creatures:
        grouped["creatures"].append(str(creature["name"]))
    grouped["creatures"].append("孢雾巡猎者")

    entry = frontmatter(type="入口", status="可用", tags=["黑曜纪元", "总览"])
    entry += "# 黑曜纪元 Vault 入口\n\n"
    entry += "这是《黑曜纪元》的 Obsidian 知识图谱入口。建议从世界底层、视觉基底、种族势力、异化生物和冲突关系五条线进入。\n\n"
    entry += "## 入口\n"
    entry += "- [[黑曜纪元]]\n- [[诸天裂灾]]\n- [[世界图谱]]\n- [[视觉基底索引]]\n- [[种族与势力索引]]\n- [[冲突关系索引]]\n- [[孢雾巡猎者]]\n"
    entry += "\n## 可视化\n"
    entry += "- [[世界关系可视化]]\n- [[黑曜纪元世界关系.canvas|黑曜纪元世界关系 Canvas]]\n"
    write_note("overview", "黑曜纪元 Vault 入口", entry)

    graph = frontmatter(type="总览图谱", status="可用", tags=["黑曜纪元", "总览", "图谱"])
    graph += "# 世界图谱\n\n"
    graph += "```mermaid\n"
    graph += "graph TD\n"
    graph += "  A[黑曜纪元] --> B[诸天裂灾]\n"
    graph += "  B --> C[高科技文明]\n"
    graph += "  B --> D[赛博工业文明]\n"
    graph += "  B --> E[克苏鲁眷族]\n"
    graph += "  B --> F[东方修行者]\n"
    graph += "  B --> G[西方魔法师]\n"
    graph += "  B --> H[北欧原始神系]\n"
    graph += "  E --> I[异化生物生态]\n"
    graph += "  I --> J[孢雾巡猎者]\n"
    graph += "  J --> K[雾绿腐生]\n"
    graph += "  J --> L[深黑不可名状]\n"
    graph += "```\n\n"
    graph += "## 使用方式\n打开 Obsidian 的 Graph View 后，从 `黑曜纪元`、`诸天裂灾`、`雾绿腐生` 或 `孢雾巡猎者` 这些节点开始扩展。\n"
    write_note("overview", "世界图谱", graph)

    index = frontmatter(type="索引", status="可用", tags=["黑曜纪元", "索引"])
    index += "# 种族与势力索引\n\n"
    labels = {"races": "种族 / 眷族", "factions": "势力 / 组织", "creatures": "异化生物"}
    for group_key in ["races", "factions", "creatures"]:
        index += f"## {labels[group_key]}\n"
        for name in sorted(set(grouped[group_key])):
            index += f"- {wikilink(name)}\n"
        index += "\n"
    write_note("overview", "种族与势力索引", index)

    base_index = frontmatter(type="索引", status="可用", tags=["黑曜纪元", "索引", "基底气质"])
    base_index += "# 视觉基底索引\n\n"
    base_index += "| 基底 | 核心气质 | 主色方向 |\n| --- | --- | --- |\n"
    for base in bases:
        base_index += f"| {wikilink(base['name'])} | {base['mood']} | {base['colors']} |\n"
    base_index += "\n## 建议\n基底决定情绪；阵营识别色和图标决定身份。卡片生成时优先引用本页，而不是重新发明色彩系统。\n"
    write_note("overview", "视觉基底索引", base_index)

    relation_index = frontmatter(type="索引", status="可用", tags=["黑曜纪元", "索引", "冲突关系"])
    relation_index += "# 冲突关系索引\n\n"
    relation_index += "| 冲突轴 | 说明 |\n| --- | --- |\n"
    for relation in relations:
        relation_index += f"| [[{relation['filename']}|{relation['axis']}]] | {relation['summary']} |\n"
    relation_index += "\n## 整理页\n- [[合并建议]]\n- [[优先主轴]]\n"
    write_note("overview", "冲突关系索引", relation_index)


def make_visualization_note() -> None:
    body = frontmatter(type="可视化", status="可用", tags=["黑曜纪元", "总览", "可视化"])
    body += "# 世界关系可视化\n\n"
    body += "本页提供静态关系图；需要可拖拽、可缩放的版本，请打开 [[黑曜纪元世界关系.canvas|黑曜纪元世界关系 Canvas]]。\n\n"
    body += "## 体系大图\n"
    body += "```mermaid\n"
    body += "graph TD\n"
    body += "  A[黑曜纪元] --> B[诸天裂灾]\n"
    body += "  B --> C[高科技文明]\n"
    body += "  B --> D[赛博工业文明]\n"
    body += "  B --> E[克苏鲁眷族]\n"
    body += "  B --> F[东方修行者]\n"
    body += "  B --> G[西方魔法师]\n"
    body += "  B --> H[北欧原始神系]\n"
    body += "  C --> C1[星序族]\n"
    body += "  C --> C2[云脑族]\n"
    body += "  C --> C3[白械天使]\n"
    body += "  D --> D1[锈裔]\n"
    body += "  D --> D2[肉件族]\n"
    body += "  D --> D3[管巢族]\n"
    body += "  E --> E1[渊眠族]\n"
    body += "  E --> E2[梦蛹族]\n"
    body += "  E --> E3[星瘤裔]\n"
    body += "  F --> F1[斩念剑庭]\n"
    body += "  F --> F2[万机符宗]\n"
    body += "  F --> F3[丹鼎肉仙门]\n"
    body += "  G --> G1[白塔秘法学院]\n"
    body += "  G --> G2[黑书会]\n"
    body += "  G --> G3[血契贵族]\n"
    body += "  H --> H1[霜巨原民]\n"
    body += "  H --> H2[狼灾血裔]\n"
    body += "  H --> H3[根下三女巫]\n"
    body += "```\n\n"
    body += "## 视觉基底映射\n"
    body += "```mermaid\n"
    body += "graph LR\n"
    body += "  A[视觉基底] --> B[钻石白]\n"
    body += "  A --> C[锈黑]\n"
    body += "  A --> D[深黑不可名状]\n"
    body += "  A --> E[雾绿腐生]\n"
    body += "  A --> F[月白青墨]\n"
    body += "  A --> G[暗金]\n"
    body += "  A --> H[血红]\n"
    body += "  A --> I[幽紫王权]\n"
    body += "  A --> J[骨白灰]\n"
    body += "  B --> B1[白械天使]\n"
    body += "  C --> C1[锈裔]\n"
    body += "  D --> D1[渊眠族]\n"
    body += "  E --> E1[孢雾巡猎者]\n"
    body += "  F --> F1[斩念剑庭]\n"
    body += "  G --> G1[白塔秘法学院]\n"
    body += "  H --> H1[狼灾血裔]\n"
    body += "  I --> I1[星序族]\n"
    body += "  J --> J1[墓歌族]\n"
    body += "```\n\n"
    body += "## 孢雾巡猎者生态图\n"
    body += "```mermaid\n"
    body += "graph TD\n"
    body += "  A[孢雾巡猎者] --> B[雾绿腐生]\n"
    body += "  A --> C[深黑不可名状]\n"
    body += "  A --> D[腐林]\n"
    body += "  A --> E[湿谷]\n"
    body += "  A --> F[废弃矿区]\n"
    body += "  A --> G[火焰]\n"
    body += "  A --> H[强光]\n"
    body += "  A --> I[干燥环境]\n"
    body += "  J[赛博工业财团] --> K[采集孢囊]\n"
    body += "  K --> A\n"
    body += "  A --> L[腐化孢子样本]\n"
    body += "  A --> M[生物武器材料]\n"
    body += "```\n"
    write_note("overview", "世界关系可视化", body)


def canvas_node(node_id: str, node_type: str, x: int, y: int, width: int, height: int, **extra) -> dict[str, object]:
    node = {"id": node_id, "type": node_type, "x": x, "y": y, "width": width, "height": height}
    node.update(extra)
    return node


def canvas_file(node_id: str, file_path: str, x: int, y: int, color: str) -> dict[str, object]:
    return canvas_node(node_id, "file", x, y, 260, 120, file=file_path, color=color)


def canvas_edge(edge_id: str, from_node: str, to_node: str, label: str = "") -> dict[str, object]:
    edge = {
        "id": edge_id,
        "fromNode": from_node,
        "fromSide": "right",
        "toNode": to_node,
        "toSide": "left",
    }
    if label:
        edge["label"] = label
    return edge


def make_canvas() -> None:
    nodes = [
        canvas_node("title", "text", -220, -660, 520, 120, text="# 黑曜纪元世界关系\n从世界底层到主轴、基底与异化生态的第一版可视化。", color="6"),
        canvas_file("world", "01_世界底层/黑曜纪元.md", -160, -420, "6"),
        canvas_file("rift", "01_世界底层/诸天裂灾.md", -160, -230, "6"),
        canvas_file("hightech", "01_世界底层/高科技文明.md", -760, 10, "2"),
        canvas_file("cyber", "01_世界底层/赛博工业文明.md", -460, 10, "1"),
        canvas_file("cthulhu", "01_世界底层/克苏鲁眷族.md", -160, 10, "4"),
        canvas_file("east", "01_世界底层/东方修行者.md", 140, 10, "5"),
        canvas_file("west", "01_世界底层/西方魔法师.md", 440, 10, "6"),
        canvas_file("norse", "01_世界底层/北欧原始神系.md", 740, 10, "3"),
        canvas_file("star-order", "02_种族/星序族.md", -900, 210, "2"),
        canvas_file("cloud-brain", "02_种族/云脑族.md", -760, 370, "2"),
        canvas_file("angel", "02_种族/白械天使.md", -620, 210, "2"),
        canvas_file("rust", "02_种族/锈裔.md", -580, 560, "1"),
        canvas_file("meatware", "02_种族/肉件族.md", -440, 720, "1"),
        canvas_file("pipe", "02_种族/管巢族.md", -300, 560, "1"),
        canvas_file("abyss", "02_种族/渊眠族.md", -260, 210, "4"),
        canvas_file("dream", "02_种族/梦蛹族.md", -120, 370, "4"),
        canvas_file("star-tumor", "02_种族/星瘤裔.md", 20, 210, "4"),
        canvas_file("sword", "03_势力组织/斩念剑庭.md", 20, 560, "5"),
        canvas_file("machine-fu", "03_势力组织/万机符宗.md", 160, 720, "5"),
        canvas_file("alchemy-flesh", "03_势力组织/丹鼎肉仙门.md", 300, 560, "5"),
        canvas_file("tower", "03_势力组织/白塔秘法学院.md", 300, 210, "6"),
        canvas_file("black-book", "03_势力组织/黑书会.md", 440, 370, "6"),
        canvas_file("blood-noble", "02_种族/血契贵族.md", 580, 210, "6"),
        canvas_file("frost", "03_势力组织/霜巨原民.md", 620, 560, "3"),
        canvas_file("wolf", "02_种族/狼灾血裔.md", 760, 720, "3"),
        canvas_file("witch", "03_势力组织/根下三女巫.md", 900, 560, "3"),
        canvas_file("creature-eco", "06_事件剧情/异化生物生态.md", -160, 950, "4"),
        canvas_file("spore", "04_异化生物/孢雾巡猎者.md", -160, 1140, "4"),
        canvas_file("green-base", "07_基底气质/雾绿腐生.md", -460, 1140, "4"),
        canvas_file("black-base", "07_基底气质/深黑不可名状.md", 140, 1140, "4"),
        canvas_file("forest", "05_地点生态/腐林.md", -460, 1320, "4"),
        canvas_file("mine", "05_地点生态/废弃矿区.md", 140, 1320, "1"),
        canvas_file("priority", "08_关系网络/优先主轴.md", 1040, 210, "6"),
        canvas_file("conflict", "00_总览/冲突关系索引.md", 1040, 390, "6"),
        canvas_file("base-index", "00_总览/视觉基底索引.md", 1040, 570, "6"),
    ]

    edges = [
        canvas_edge("e-world-rift", "world", "rift", "源头"),
        canvas_edge("e-rift-hightech", "rift", "hightech"),
        canvas_edge("e-rift-cyber", "rift", "cyber"),
        canvas_edge("e-rift-cthulhu", "rift", "cthulhu"),
        canvas_edge("e-rift-east", "rift", "east"),
        canvas_edge("e-rift-west", "rift", "west"),
        canvas_edge("e-rift-norse", "rift", "norse"),
        canvas_edge("e-hightech-star", "hightech", "star-order"),
        canvas_edge("e-hightech-cloud", "hightech", "cloud-brain"),
        canvas_edge("e-hightech-angel", "hightech", "angel"),
        canvas_edge("e-cyber-rust", "cyber", "rust"),
        canvas_edge("e-cyber-meat", "cyber", "meatware"),
        canvas_edge("e-cyber-pipe", "cyber", "pipe"),
        canvas_edge("e-cthulhu-abyss", "cthulhu", "abyss"),
        canvas_edge("e-cthulhu-dream", "cthulhu", "dream"),
        canvas_edge("e-cthulhu-tumor", "cthulhu", "star-tumor"),
        canvas_edge("e-east-sword", "east", "sword"),
        canvas_edge("e-east-machine", "east", "machine-fu"),
        canvas_edge("e-east-flesh", "east", "alchemy-flesh"),
        canvas_edge("e-west-tower", "west", "tower"),
        canvas_edge("e-west-book", "west", "black-book"),
        canvas_edge("e-west-blood", "west", "blood-noble"),
        canvas_edge("e-norse-frost", "norse", "frost"),
        canvas_edge("e-norse-wolf", "norse", "wolf"),
        canvas_edge("e-norse-witch", "norse", "witch"),
        canvas_edge("e-rift-eco", "rift", "creature-eco", "异化生态"),
        canvas_edge("e-eco-spore", "creature-eco", "spore"),
        canvas_edge("e-green-spore", "green-base", "spore", "基底"),
        canvas_edge("e-black-spore", "black-base", "spore", "基底"),
        canvas_edge("e-forest-spore", "forest", "spore", "栖息"),
        canvas_edge("e-mine-spore", "mine", "spore", "栖息"),
        canvas_edge("e-priority", "priority", "world"),
        canvas_edge("e-conflict", "conflict", "world"),
        canvas_edge("e-base-index", "base-index", "world"),
    ]

    canvas = {"nodes": nodes, "edges": edges}
    path = ROOT / FOLDERS["overview"] / "黑曜纪元世界关系.canvas"
    path.write_text(json.dumps(canvas, ensure_ascii=False, indent=2), encoding="utf-8")


def make_templates() -> None:
    templates = {
        "种族模板": """---
type: 种族
base:
status: 草案
tags:
  - 黑曜纪元
  - 种族
---

# {{title}}

## 关系
- 基底：
- 所属体系：
- 敌对 / 合作：

## 外观与气质

## 核心能力

## 代价与弱点

## 社会结构

## 剧情用途
""",
        "势力组织模板": """---
type: 势力组织
base:
status: 草案
tags:
  - 黑曜纪元
  - 势力组织
---

# {{title}}

## 关系
- 所属体系：
- 控制资源：
- 敌对 / 合作：

## 权力结构

## 信条或目标

## 代表角色

## 剧情用途
""",
        "异化生物模板": """---
type: 异化生物
base:
habitat:
threat:
weakness:
status: 草案
tags:
  - 黑曜纪元
  - 异化生物
---

# {{title}}

## 关系
- 生态位：
- 栖息地：
- 可采集资源：

## 外观

## 行为

## 能力

## 弱点

## 遭遇事件
""",
        "基底气质模板": """---
type: 基底气质
status: 草案
tags:
  - 黑曜纪元
  - 基底气质
---

# {{title}}

## 核心气质

## 色彩方向

## 适配对象

## 禁忌
""",
    }
    for name, text in templates.items():
        write_note("templates", name, text)


def make_obsidian_config() -> None:
    (ROOT / ".obsidian" / "app.json").write_text(
        '{\n  "attachmentFolderPath": "09_素材与图片",\n  "alwaysUpdateLinks": true,\n  "newFileLocation": "current"\n}\n',
        encoding="utf-8",
    )
    (ROOT / ".obsidian" / "appearance.json").write_text(
        '{\n  "theme": "obsidian",\n  "accentColor": "#2f8f77"\n}\n',
        encoding="utf-8",
    )


def stub_folder_for(target: str) -> str:
    if re.match(r"^\d+\.", target) or any(word in target for word in ["生态位", "异化生物"]):
        return "plot"
    if any(word in target for word in ["腐林", "湿谷", "矿区", "深海", "森林", "管道", "废液", "高空", "空间"]):
        return "places"
    if any(word in target for word in ["文明", "体系", "神系", "修行者", "修士", "魔法师", "眷族"]):
        return "world"
    if any(word in target for word in ["财团", "学院", "宗", "庭", "会", "族", "民", "修女"]):
        return "factions"
    return "relations"


def make_missing_link_stubs() -> None:
    notes = {path.stem for path in ROOT.rglob("*.md")}
    existing_files = {path.name for path in ROOT.rglob("*") if path.is_file()}
    targets: set[str] = set()
    for path in ROOT.rglob("*.md"):
        text = path.read_text(encoding="utf-8")
        for match in re.findall(r"\[\[([^\]]+)\]\]", text):
            target = match.split("|", 1)[0].split("#", 1)[0].strip()
            if target:
                targets.add(target)

    for target in sorted(targets - notes):
        if target in existing_files:
            continue
        body = frontmatter(type="图谱节点", status="占位", tags=["黑曜纪元", "图谱节点"])
        body += f"# {target}\n\n"
        body += "此节点由导入器自动创建，用于补齐 Obsidian 图谱中的核心关系。后续可以把它扩写成完整设定条目。\n\n"
        body += f"## 关联\n- {wikilink('黑曜纪元')}\n"
        write_note(stub_folder_for(target), target, body)


def make_design_note() -> None:
    body = """# Obsidian 知识图谱设计

## 目标

将《黑曜纪元》的种族、势力、异化生物、基底气质、地点生态和冲突关系拆成可双链的 Obsidian 条目。第一版优先保证写作可读、图谱可看、后续可扩展，不引入重型数据库。

## 结构

- `00_总览`：入口、索引和总图。
- `01_世界底层`：世界名、诸天裂灾等母设定。
- `02_种族`：种族、眷族、血裔和造物种族。
- `03_势力组织`：宗门、学院、教会、财团、部族和群体。
- `04_异化生物`：生态怪物、野外遭遇和图像设定。
- `05_地点生态`：腐林、湿谷、矿区等地点节点。
- `06_事件剧情`：事件、剧情线和生态用途。
- `07_基底气质`：视觉基底与色彩系统。
- `08_关系网络`：冲突轴、合并建议和主轴规划。
- `09_素材与图片`：图片、参考图和卡片素材。

## 数据约定

每个条目使用 YAML 前置字段保存 `type`、`base`、`status`、`tags` 等轻量结构化信息；正文使用 Obsidian 双链表达关系。这样既能用原生 Graph View 查看关系，也能在未来接入 Dataview 或自动导出脚本。

## 第一版范围

本次导入 DOCX 中的基底气质、种族、势力、异化生物生态、冲突轴、合并建议和优先主轴，并新增两张图对应的 `孢雾巡猎者` 条目。
"""
    path = ROOT / FOLDERS["docs"] / "2026-06-04-obsidian-knowledge-graph-design.md"
    path.write_text(body, encoding="utf-8")


def main() -> None:
    mkdirs()
    doc = Document(SOURCE_DOC)
    copy_assets()
    make_obsidian_config()
    make_design_note()
    make_world_notes(doc, doc.tables[1])
    make_hub_notes()
    bases = make_base_notes(doc.tables[2])

    source_sections = [
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "3. 通用诡秘种族库",
        "4. 高科技风种族",
        "4. 高科技风种族",
        "4. 高科技风种族",
        "4. 高科技风种族",
        "4. 高科技风种族",
        "4. 高科技风种族",
        "5. 赛博朋克工业风种族",
        "5. 赛博朋克工业风种族",
        "5. 赛博朋克工业风种族",
        "5. 赛博朋克工业风种族",
        "5. 赛博朋克工业风种族",
        "5. 赛博朋克工业风种族",
        "5. 赛博朋克工业风种族",
        "5. 赛博朋克工业风种族",
        "6. 克苏鲁眷族",
        "6. 克苏鲁眷族",
        "6. 克苏鲁眷族",
        "6. 克苏鲁眷族",
        "6. 克苏鲁眷族",
        "7. 东方修行群体",
        "7. 东方修行群体",
        "7. 东方修行群体",
        "7. 东方修行群体",
        "7. 东方修行群体",
        "7. 东方修行群体",
        "8. 西方魔法群体",
        "8. 西方魔法群体",
        "8. 西方魔法群体",
        "8. 西方魔法群体",
        "8. 西方魔法群体",
        "8. 西方魔法群体",
        "9. 北欧原始神系与部族",
        "9. 北欧原始神系与部族",
        "9. 北欧原始神系与部族",
        "9. 北欧原始神系与部族",
        "9. 北欧原始神系与部族",
        "9. 北欧原始神系与部族",
    ]

    entities = []
    for table, section in zip(doc.tables[3:50], source_sections, strict=True):
        entity = extract_card_table(table)
        entities.append(entity)
        write_note(classify_folder(str(entity["category"])), str(entity["name"]), make_entity_note(entity, section))

    creatures = make_creature_ecology_notes(doc.tables[50:55])
    make_spore_hunter_note()
    relations = make_relation_notes(doc.tables[55])
    make_merge_and_priority_notes(doc.tables[56], doc.tables[57])
    make_templates()
    make_indexes(entities, bases, creatures, relations)
    make_visualization_note()
    make_canvas()
    make_missing_link_stubs()


if __name__ == "__main__":
    main()
